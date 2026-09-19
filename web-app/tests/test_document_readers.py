from io import BytesIO
from pathlib import Path
import unittest

from docx import Document
from reportlab.pdfgen.canvas import Canvas
from src.extract import extract_attachment_with_evidence, ReviewRequired
from test_workflows import SI_DOCUMENT


class BytesInbox:
    def __init__(self, raw): self.raw = raw
    def read_bytes(self, path): return self.raw


class ReaderTests(unittest.TestCase):
    def test_pdf_second_page_evidence(self):
        out = BytesIO()
        canvas = Canvas(out)
        canvas.drawString(40, 780, 'SHIPPING INSTRUCTION')
        canvas.showPage()
        y = 780
        for line in SI_DOCUMENT.splitlines()[2:]:
            canvas.drawString(40, y, line)
            y -= 20
        canvas.save()
        fields, evidence = extract_attachment_with_evidence(BytesInbox(out.getvalue()), 'test_SI.pdf', 'SI')
        self.assertEqual(fields['container_count'], '1')
        self.assertEqual(evidence['shipper']['page'], 2)

    def test_word_table_and_label_units(self):
        doc = Document()
        doc.add_paragraph('SHIPPING INSTRUCTION')
        table = doc.add_table(rows=0, cols=2)
        for line in SI_DOCUMENT.splitlines()[2:]:
            label, value = line.split(':', 1)
            cells = table.add_row().cells
            cells[0].text = 'Gross Weight (KG)' if label == 'Gross Weight' else label
            cells[1].text = '1000' if label == 'Gross Weight' else value.strip()
        out = BytesIO(); doc.save(out)
        fields, evidence = extract_attachment_with_evidence(BytesInbox(out.getvalue()), 'test_SI.docx', 'SI')
        self.assertEqual(fields['gross_weight_kg'], '1000 KG')
        self.assertIsNone(evidence['shipper']['page'])

    def test_blank_pdf_goes_to_review(self):
        out = BytesIO(); canvas = Canvas(out); canvas.showPage(); canvas.save()
        with self.assertRaises(ReviewRequired) as context:
            extract_attachment_with_evidence(BytesInbox(out.getvalue()), 'scan_SI.pdf', 'SI')
        self.assertEqual(context.exception.internal_reason, 'document_read_failed')

    def test_corrupt_docx_goes_to_review(self):
        with self.assertRaises(ReviewRequired):
            extract_attachment_with_evidence(BytesInbox(b'broken'), 'test_SI.docx', 'SI')
