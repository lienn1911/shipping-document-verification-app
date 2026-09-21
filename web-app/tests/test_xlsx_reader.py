"""XLSX attachments: reading, cross-format comparison, and failure handling. No network."""

from io import BytesIO
import logging
from pathlib import Path
import sys
import unittest

logging.getLogger("pypdf").setLevel(logging.ERROR)  # the dataset contains deliberately corrupt PDFs

WEB_APP = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(WEB_APP))

from docx import Document  # noqa: E402
from openpyxl import Workbook  # noqa: E402

from src.compare import compare_documents  # noqa: E402
from src.doctype import BL, SI, detect_document_type  # noqa: E402
from src.document_readers import read_document  # noqa: E402
from src.extract import TEXT_READERS, ReviewRequired, extract_attachment  # noqa: E402
from src.service import process_dataset  # noqa: E402

BUNDLE = WEB_APP.parent / "local-data" / "participant-bundle"


def workbook_bytes(rows, sheet="S.I.") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for row in rows:
        ws.append(row)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


SI_ROWS = [
    ["ACME TRADING PTE LTD", None],
    [None, None],
    ["BL INSTRUCTION", "3154303911"],
    ["SHIPPER", "ACME TRADING PTE LTD | 80 RAFFLES PLACE; SINGAPORE 048624"],
    ["Consignee (Non-Negotiable)", "BUYER LLC | P.O. BOX 5069; DUBAI, UAE"],
    ["NOTIFY PARTY", "BUYER LLC | P.O. BOX 5069; DUBAI, UAE"],
    ["Port of Loading (POL)", "SINGAPORE"],
    ["Port of Discharge", "KARACHI, PAKISTAN"],
    ["No. of Containers or Packages", "12 x 20'GP"],
    ["GROSS WEIGHT", 243588],
    ["BOOKING NO.", "PSGSE8148932"],
]


class FakeInbox:
    def __init__(self, files):
        self.files = files

    def read_bytes(self, path):
        return self.files[path]


def bl_docx_bytes(consignee_lines="BUYER LLC\nP.O. BOX 5069\nDUBAI, UAE", weight="243,588") -> bytes:
    doc = Document()
    doc.add_paragraph("BILL OF LADING (DRAFT)")
    table = doc.add_table(rows=0, cols=2)
    for label, value in (
        ("Shipper (Principal or Seller)", "ACME TRADING PTE LTD\n80 RAFFLES PLACE\nSINGAPORE 048624"),
        ("Consignee", consignee_lines),
        ("Notify", "BUYER LLC\nP.O. BOX 5069\nDUBAI, UAE"),
        ("Load Port", "SINGAPORE"),
        ("POD", "KARACHI, PAKISTAN"),
        ("Total Containers", "12 x 20'GP"),
        ("Gross Wt (kgs)", weight),
    ):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, value
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


class ReaderTests(unittest.TestCase):
    def test_rows_are_label_value_with_sheet_and_row_numbers(self):
        rows = read_document(workbook_bytes(SI_ROWS), "x_SI.xlsx")
        self.assertEqual(rows[0], ("ACME TRADING PTE LTD", 1, 1))  # single-cell letterhead row
        self.assertIn(("Port of Discharge: KARACHI, PAKISTAN", 1, 8), rows)
        self.assertNotIn(None, [r[0] for r in rows])  # empty rows are skipped, not emitted

    def test_bare_weight_number_gets_a_unit_and_is_never_doubled(self):
        for cell in (243588, 243588.0, "243,588", "243588 KG"):
            rows = dict((line.split(": ")[0], line) for line, _, _ in read_document(workbook_bytes([["Gross Weight", cell]]), "a.xlsx"))
            self.assertEqual(rows["Gross Weight"].count("KG"), 1, cell)
        self.assertIn("243588 KG", [line.split(": ")[1] for line, _, _ in read_document(workbook_bytes([["GROSS WEIGHT", 243588]]), "a.xlsx")])
        # a non-weight number is left alone
        self.assertEqual(read_document(workbook_bytes([["HS CODE", 48109200]]), "a.xlsx")[0][0], "HS CODE: 48109200")

    def test_party_separators_become_plain_words(self):
        rows = read_document(workbook_bytes([["SHIPPER", "APRIL FINE | ON BEHALF OF VITAL; 77 ROAD, #21-01; SINGAPORE 068896"]]), "a.xlsx")
        self.assertEqual(rows[0][0], "SHIPPER: APRIL FINE ON BEHALF OF VITAL 77 ROAD, #21-01 SINGAPORE 068896")

    def test_type_detection_uses_the_title_row(self):
        text = TEXT_READERS[".xlsx"](workbook_bytes(SI_ROWS))
        self.assertEqual(detect_document_type(text).type, SI)
        bl_rows = [["ACME TRADING PTE LTD", None], ["BILL OF LADING", "3154303911"], ["Shipper", "ACME"]]
        self.assertEqual(detect_document_type(TEXT_READERS[".xlsx"](workbook_bytes(bl_rows, "BL"))).type, BL)


class ExtractionTests(unittest.TestCase):
    def fields(self, path, raw, role):
        return extract_attachment(FakeInbox({path: raw}), path, role)

    def test_xlsx_si_and_docx_bl_compare_without_false_alarms(self):
        si = self.fields("e_SI.xlsx", workbook_bytes(SI_ROWS), "SI")
        bl = self.fields("e_BL.docx", bl_docx_bytes(), "BL")
        self.assertEqual(compare_documents(si, bl)["status"], "OK")

    def test_real_differences_are_still_caught_across_formats(self):
        si = self.fields("e_SI.xlsx", workbook_bytes(SI_ROWS), "SI")
        bl = self.fields("e_BL.docx", bl_docx_bytes(consignee_lines="OTHER CO LTD\nP.O. BOX 1\nDUBAI, UAE", weight="243,000"), "BL")
        result = compare_documents(si, bl)
        self.assertEqual(sorted(m["field"] for m in result["mismatches"]), ["consignee", "gross_weight_kg"])

    def test_corrupt_or_empty_workbooks_are_reported_not_guessed(self):
        for label, raw in (("corrupt", b"this is not an xlsx"), ("empty", workbook_bytes([]))):
            with self.assertRaises(ReviewRequired, msg=label) as ctx:
                self.fields("e_SI.xlsx", raw, "SI")
            self.assertEqual((ctx.exception.review_reason, ctx.exception.internal_reason), ("unreadable", "document_read_failed"))

    def test_a_bill_of_lading_in_the_si_slot_is_rejected(self):
        bl_rows = [["ACME", None], ["BILL OF LADING", "1"], ["Shipper", "ACME"]]
        with self.assertRaises(ReviewRequired) as ctx:
            self.fields("e_SI.xlsx", workbook_bytes(bl_rows, "BL"), "SI")
        self.assertEqual(ctx.exception.review_reason, "wrong_doc_type")


@unittest.skipUnless((BUNDLE / "inbox").is_dir(), "participant bundle not present")
class DatasetTests(unittest.TestCase):
    def test_no_email_is_left_unsupported(self):
        artifacts = process_dataset(BUNDLE)
        reasons = [d.get("internal_reason") for d in artifacts.internal_results.values()]
        self.assertNotIn("unsupported_attachment_format", reasons)
        xlsx_emails = [e["email_id"] for e in artifacts.emails if any(p.endswith(".xlsx") for p in e["attachments"])]
        self.assertEqual(len(xlsx_emails), 15)
        for email_id in xlsx_emails:
            self.assertIn(artifacts.submission[email_id]["status"], {"OK", "MISMATCH"}, email_id)


if __name__ == "__main__":
    unittest.main()
