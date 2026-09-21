"""Text-layer PDF, paragraph/table DOCX and label/value XLSX readers. No OCR or guessed values."""
from io import BytesIO
from pathlib import Path
import re

_XLSX_MAX_SHEETS = 5
_XLSX_MAX_ROWS = 500


def _xlsx_cell_text(value) -> str:
    """Cell content as one line of text. Numbers lose a spurious ".0"; line breaks become spaces."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return re.sub(r"\s+", " ", str(value)).strip()


def _read_xlsx(raw: bytes):
    """Rows of "Label: value" (first cell is the label, the rest is the value).

    * Party layouts differ between formats ("NAME | ADDRESS; LINE" in a sheet, one line per part in Word),
      so "|" and ";" separators become spaces. That keeps an XLSX SI comparable with a DOCX or PDF BL.
    * A weight cell holding only a number gets " KG" (the field is defined in kilograms and the sheets
      carry no unit), so it goes through the same weight normalisation as every other format.
    * Only cached cell values are read: formulas, macros and external links are never executed.
    """
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:  # BadZipFile, InvalidFileException, KeyError ... all mean "not a usable workbook"
        raise ValueError("Excel workbook could not be opened (corrupt or not a real .xlsx file)") from exc
    rows = []
    try:
        for sheet_no, sheet in enumerate(workbook.worksheets[:_XLSX_MAX_SHEETS], 1):
            for row_no, cells in enumerate(sheet.iter_rows(max_row=_XLSX_MAX_ROWS, values_only=True), 1):
                parts = [text for text in (_xlsx_cell_text(c) for c in cells) if text]
                if not parts:
                    continue
                label, values = parts[0], parts[1:]
                if not values:
                    rows.append((label, sheet_no, row_no))
                    continue
                value = re.sub(r"\s*[|;]\s*", " ", " ".join(values)).strip()
                if re.search(r"weight|\bwt\b", label, re.I) and re.fullmatch(r"[\d,.\s]+", value):
                    value += " KG"
                rows.append((f"{label}: {value}", sheet_no, row_no))
    finally:
        workbook.close()
    if not rows:
        raise ValueError("Excel workbook contains no readable cells")
    return rows


def read_document(raw: bytes, path: str):
    suffix = Path(path).suffix.lower()
    if suffix == ".txt":
        return [(line, 1, n) for n, line in enumerate(raw.decode("utf-8").splitlines(), 1)]
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(raw))
        if reader.is_encrypted:
            raise ValueError("Encrypted PDF; provide an unlocked copy")
        if len(reader.pages) > 100:
            raise ValueError("PDF exceeds the 100-page reading limit")
        rows = []
        for page_no, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            if not text.strip():
                raise ValueError(f"PDF page {page_no} has no readable text layer; scanned pages require OCR/manual review")
            rows.extend((line, page_no, n) for n, line in enumerate(text.splitlines(), 1))
        return rows
    if suffix == ".docx":
        from docx import Document
        from docx.table import Table
        doc = Document(BytesIO(raw))
        rows = []
        for block in doc.iter_inner_content():
            if isinstance(block, Table):
                for row in block.rows:
                    cells = list(dict.fromkeys(c.text.strip() for c in row.cells))
                    text = cells[0] + ": " + " ".join(cells[1:]) if len(cells) > 1 else " ".join(cells)
                    rows.append((text.replace("\n", " "), None, len(rows) + 1))
            else:
                rows.extend((line, None, len(rows) + n) for n, line in enumerate(block.text.splitlines(), 1))
        if not any(line.strip() for line, _, _ in rows):
            raise ValueError("Word document contains no readable paragraphs or tables")
        return rows
    if suffix == ".xlsx":
        return _read_xlsx(raw)
    raise ValueError(f"Unsupported document format: {suffix}")


def document_text(raw: bytes, path: str) -> str:
    return "\n".join(line for line, _, _ in read_document(raw, path))
